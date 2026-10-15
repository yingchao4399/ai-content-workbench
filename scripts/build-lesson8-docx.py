from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'LESSON8-HOMEWORK.md'
OUTPUT = ROOT / 'artifacts' / 'lesson8-homework-report.docx'

NAVY = '17365D'
LIGHT_BLUE = 'D9EAF7'
LIGHT_GRAY = 'F4F6F8'
BORDER = 'C8D1DC'
TEXT = '20262E'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=BORDER, size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_font(run, size=10.5, bold=False, color=TEXT, name='Microsoft YaHei'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_inline(paragraph, value, base_size=10.5):
    # Keep links readable as ordinary text; Feishu and Word both preserve the URL.
    tokens = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', value)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_font(run, base_size, True)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            set_font(run, 9.2, False, '303840', 'Consolas')
        else:
            run = paragraph.add_run(token)
            set_font(run, base_size)

def clean_heading(text):
    text = re.sub(r'^[0-9]+[.]\s*', lambda m: m.group(0).replace('.', ''), text)
    return text.replace('｜', ' ').replace('：', ' ').replace(':', ' ')

def add_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=0, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, NAVY)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            run_text = re.sub(r'\s+', ' ', value).strip()
            add_inline(p, run_text, 9.1 if row_idx else 9.1)
            for run in p.runs:
                if row_idx == 0:
                    set_font(run, 9.1, True, 'FFFFFF')
        
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25
    for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 11.5)]:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('000000')
        style.paragraph_format.space_before = Pt(13 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line or line == '---':
            i += 1
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, line[2:].replace('：', ' '), 20)
            p.paragraph_format.space_after = Pt(10)
        elif line.startswith('## '):
            p = doc.add_paragraph(clean_heading(line[3:]), style='Heading 1')
        elif line.startswith('### '):
            p = doc.add_paragraph(clean_heading(line[4:]), style='Heading 2')
        elif line.startswith('!['):
            m = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
            if m:
                image_path = ROOT / m.group(2)
                if image_path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_with_next = True
                    p.add_run().add_picture(str(image_path), width=Cm(15.4))
                    cap = doc.add_paragraph(m.group(1))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(8)
                    for run in cap.runs:
                        set_font(run, 9, False, '5A6570')
        elif line.startswith('| '):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [re.split(r'\s*\|\s*', row.strip('| ')) for row in table_lines]
            parsed = [row for row in parsed if not all(re.fullmatch(r'-+', cell.strip()) for cell in row)]
            add_table(doc, parsed)
            continue
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline(p, line[2:])
        elif re.match(r'^\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number')
            add_inline(p, re.sub(r'^\d+\.\s+', '', line))
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('第八课作业 · EasyInput 剪贴板速记 · 2026-10-15')
    set_font(r, 8, False, '66717D')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__':
    main()
