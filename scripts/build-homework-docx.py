"""Build the fourth-lesson submission with embedded local evidence screenshots."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "HOMEWORK-SUBMISSION.md"
OUTPUT = ROOT / "artifacts" / "第四课作业 AI 内容创作工作台实践报告.docx"


def add_inline(paragraph, value):
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(36, 43, 52)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28
    for name, size in [("Title", 19), ("Heading 1", 14)]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(28, 55, 76)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                path = ROOT / match.group(2)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(path), width=Cm(15.2))
        elif line.startswith("图 "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(line)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(91, 101, 112)
        elif line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [re.split(r"\s*\|\s*", x.strip("| ")) for x in table_lines]
            rows = [r for r in rows if not all(re.fullmatch(r"-+", c) for c in r)]
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    add_inline(cell.paragraphs[0], value)
            i -= 1
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_inline(doc.add_paragraph(), line)
        i += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第四课实践报告 · 2026-10-14 · 待作者审稿").font.size = Pt(8)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
